"""
Import unico da planilha de inventario para a tabela `equipment` do SQLite.

Rode UMA vez, antes de subir a versao que nao usa mais Excel:

    python import_inventory.py                 # local (venv com openpyxl)
    docker compose run --rm pingador python import_inventory.py

A partir dai o app usa apenas o banco (ver db.py / storage.py). Este script
e independente do resto do codigo (traz sua propria leitura de planilha) e
pode ser apagado depois do import.

Idempotente: rodar de novo nao duplica (IPs ja presentes sao ignorados).

Preserva o historico: quando o IP ja tem snapshot em `equipment_state`, o
equipamento e inserido com o MESMO id, mantendo `ping_samples` / `events` /
`equipment_state` ligados.
"""
import ipaddress
import os
import sys
import unicodedata

import openpyxl

import db

EXCEL_PATH = os.environ.get(
    "PINGADOR_EXCEL_PATH",
    os.path.join(os.path.dirname(os.path.abspath(__file__)), "Inventario IP.xlsx"),
)
DEFAULT_FREQUENCY = int(os.environ.get("PINGADOR_EXCEL_DEFAULT_FREQUENCY", "15"))
IGNORED_SHEETS = {"geral", "pesquisa"}


def _normalize(text) -> str:
    if text is None:
        return ""
    text = str(text).strip().lower()
    return unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")


def _find_columns(header_row) -> dict:
    columns = {}
    for idx, cell in enumerate(header_row):
        key = _normalize(cell)
        if key == "ip":
            columns["ip"] = idx
        elif key == "tipo":
            columns["tipo"] = idx
        elif key == "modelo":
            columns["modelo"] = idx
        elif key == "usuario":
            columns["usuario"] = idx
        elif key in ("observacao", "observacoes"):
            columns["observacao"] = idx
    return columns


def _cell(row, columns, name):
    idx = columns.get(name)
    if idx is None or idx >= len(row):
        return None
    value = row[idx]
    return str(value).strip() if value not in (None, "") else None


def read_inventory(path: str = None) -> list[dict]:
    """Le a planilha e retorna uma lista de equipamentos, um por IP."""
    path = path or EXCEL_PATH
    entries_by_ip = {}
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        for sheet_name in wb.sheetnames:
            if _normalize(sheet_name) in IGNORED_SHEETS:
                continue
            ws = wb[sheet_name]
            rows = ws.iter_rows(values_only=True)
            try:
                header_row = next(rows)
            except StopIteration:
                continue
            columns = _find_columns(header_row)
            if "ip" not in columns:
                continue

            for row in rows:
                ip_raw = _cell(row, columns, "ip")
                if not ip_raw:
                    continue
                try:
                    ipaddress.ip_address(ip_raw)
                except ValueError:
                    continue
                if ip_raw in entries_by_ip:
                    continue

                tipo = _cell(row, columns, "tipo")
                modelo = _cell(row, columns, "modelo")
                usuario = _cell(row, columns, "usuario")
                observacao = _cell(row, columns, "observacao")

                name_parts = [p for p in (tipo, modelo) if p]
                name = " - ".join(name_parts) if name_parts else ip_raw

                desc_parts = []
                if observacao:
                    desc_parts.append(observacao)
                if usuario:
                    desc_parts.append(f"Usuario: {usuario}")

                entries_by_ip[ip_raw] = {
                    "ip": ip_raw,
                    "name": name,
                    "description": " | ".join(desc_parts),
                    "category": sheet_name,
                    "frequency": DEFAULT_FREQUENCY,
                }
    finally:
        wb.close()
    return list(entries_by_ip.values())


def main() -> int:
    if not os.path.exists(EXCEL_PATH):
        print(f"Planilha nao encontrada: {EXCEL_PATH}", file=sys.stderr)
        return 1

    entries = read_inventory()
    existing_ips = {e["ip"] for e in db.list_equipment()}
    state_by_ip = {}
    for s in db.list_states():
        ip = s.get("ip")
        if ip and ip not in state_by_ip:
            state_by_ip[ip] = s

    imported = adopted = skipped = 0
    for entry in entries:
        ip = entry["ip"]
        if ip in existing_ips:
            skipped += 1
            continue
        adopt_id = None
        st = state_by_ip.get(ip)
        if st and st.get("equipment_id") is not None:
            adopt_id = int(st["equipment_id"])
        try:
            db.insert_equipment(
                ip=ip,
                name=entry["name"],
                description=entry["description"],
                frequency=entry["frequency"],
                category=entry["category"],
                source="excel",
                monitoring_active=True,
                equipment_id=adopt_id,
            )
        except db.EquipmentIPConflict:
            # id historico ja ocupado por outro IP: cai no AUTOINCREMENT
            db.insert_equipment(
                ip=ip,
                name=entry["name"],
                description=entry["description"],
                frequency=entry["frequency"],
                category=entry["category"],
                source="excel",
                monitoring_active=True,
            )
            adopt_id = None
        imported += 1
        if adopt_id is not None:
            adopted += 1

    print(
        f"Planilha: {len(entries)} IPs. "
        f"Importados: {imported} (adotaram id historico: {adopted}). "
        f"Ja existiam: {skipped}."
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
