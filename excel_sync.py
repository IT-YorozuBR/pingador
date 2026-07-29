"""
Sincronizacao do inventario de IPs a partir da planilha Excel.

Le periodicamente o arquivo de inventario (uma aba por categoria/base de
equipamentos: Servidores, Access Points, Cameras, Impressoras, etc.) e
mantem os equipamentos cadastrados no sistema em sincronia:

- Adiciona os IPs novos encontrados em qualquer aba.
- Atualiza nome/descricao/categoria quando mudam na planilha.
- Remove os equipamentos que somem da planilha.

Equipamentos cadastrados manualmente pela tela (source="manual") nunca sao
tocados por essa sincronizacao, mesmo que o mesmo IP exista na planilha.

As abas "Geral" e "Pesquisa" sao ignoradas: sao visoes de apoio (mapeamento
para listas suspensas / busca), nao uma lista real de equipamentos por
categoria.
"""
import ipaddress
import os
import threading
import unicodedata

import openpyxl

from storage import store

EXCEL_PATH = os.environ.get(
    "PINGADOR_EXCEL_PATH",
    os.path.join(os.path.dirname(os.path.abspath(__file__)), "Inventario IP.xlsx"),
)
SYNC_INTERVAL_SECONDS = int(os.environ.get("PINGADOR_EXCEL_SYNC_INTERVAL", "60"))
DEFAULT_FREQUENCY = int(os.environ.get("PINGADOR_EXCEL_DEFAULT_FREQUENCY", "15"))

IGNORED_SHEETS = {"geral", "pesquisa"}

_stop_event = threading.Event()
_watcher_thread = None
_last_mtime = None


def _normalize(text) -> str:
    if text is None:
        return ""
    text = str(text).strip().lower()
    return unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")


def _find_columns(header_row) -> dict:
    columns = {}
    for idx, cell in enumerate(header_row):
        key = _normalize(cell)
        if key == "ip":
            columns["ip"] = idx
        elif key == "tipo":
            columns["tipo"] = idx
        elif key == "modelo":
            columns["modelo"] = idx
        elif key == "usuario":
            columns["usuario"] = idx
        elif key in ("observacao", "observacoes"):
            columns["observacao"] = idx
    return columns


def _cell(row, columns, name):
    idx = columns.get(name)
    if idx is None or idx >= len(row):
        return None
    value = row[idx]
    return str(value).strip() if value not in (None, "") else None


def read_inventory(path: str = None) -> list[dict]:
    """
    Le a planilha e retorna uma lista de equipamentos, um por IP.
    Em caso de IP duplicado entre abas, a primeira aba em que ele aparece
    prevalece (as abas sao lidas na ordem em que aparecem no arquivo).
    """
    path = path or EXCEL_PATH
    entries_by_ip = {}
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        for sheet_name in wb.sheetnames:
            if _normalize(sheet_name) in IGNORED_SHEETS:
                continue
            ws = wb[sheet_name]
            rows = ws.iter_rows(values_only=True)
            try:
                header_row = next(rows)
            except StopIteration:
                continue
            columns = _find_columns(header_row)
            if "ip" not in columns:
                continue

            for row in rows:
                ip_raw = _cell(row, columns, "ip")
                if not ip_raw:
                    continue
                try:
                    ipaddress.ip_address(ip_raw)
                except ValueError:
                    continue
                if ip_raw in entries_by_ip:
                    continue

                tipo = _cell(row, columns, "tipo")
                modelo = _cell(row, columns, "modelo")
                usuario = _cell(row, columns, "usuario")
                observacao = _cell(row, columns, "observacao")

                name_parts = [p for p in (tipo, modelo) if p]
                name = " - ".join(name_parts) if name_parts else ip_raw

                desc_parts = []
                if observacao:
                    desc_parts.append(observacao)
                if usuario:
                    desc_parts.append(f"Usuario: {usuario}")

                entries_by_ip[ip_raw] = {
                    "ip": ip_raw,
                    "name": name,
                    "description": " | ".join(desc_parts),
                    "category": sheet_name,
                    "frequency": DEFAULT_FREQUENCY,
                }
    finally:
        wb.close()
    return list(entries_by_ip.values())


def sync_now(path: str = None) -> dict:
    entries = read_inventory(path)
    return store.sync_from_excel(entries)


def _watch_loop():
    global _last_mtime
    while not _stop_event.is_set():
        try:
            mtime = os.path.getmtime(EXCEL_PATH)
            if mtime != _last_mtime:
                _last_mtime = mtime
                sync_now()
        except OSError:
            pass  # arquivo ausente/indisponivel momentaneamente
        _stop_event.wait(SYNC_INTERVAL_SECONDS)


def start_excel_watcher():
    global _watcher_thread
    if _watcher_thread is None or not _watcher_thread.is_alive():
        _stop_event.clear()
        _watcher_thread = threading.Thread(target=_watch_loop, daemon=True)
        _watcher_thread.start()


def stop_excel_watcher():
    _stop_event.set()
